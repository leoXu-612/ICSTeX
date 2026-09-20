from copy import deepcopy
from dataclasses import replace
import hashlib
import io
import re
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from app.core.blocks.source_registry import SourceRecord
from app.core.blocks.source_repair import (
    ColumnMapping, ImportOptions, capture_source, map_source_tables, parse_source,
    verify_capture,
)
from app.core.blocks.table_import import read_csv_text


class SourceRepairTests(TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.original = b"Key,Value\nA,10\nB,20\n"
        (self.root / "base.csv").write_bytes(self.original)
        (self.root / "data.csv").write_bytes(b"Value,Key\n21,B\n11,A\n30,C\n")
        self.record = SourceRecord("source", "csv", "data.csv", hashlib.sha256(self.original).hexdigest())

    def test_capture_requires_real_baseline_and_detects_later_source_changes(self):
        base = capture_source(self.root, "base.csv", expected_sha=self.record.baseSha256)
        remote = capture_source(self.root, "data.csv")
        self.assertEqual(base.payload, self.original)
        self.assertTrue(verify_capture(remote))
        (self.root / "data.csv").write_bytes(b"changed")
        self.assertFalse(verify_capture(remote))
        with self.assertRaises(ValueError):
            capture_source(self.root, "data.csv", expected_sha=self.record.baseSha256)

    def test_capture_rejects_links_internal_paths_and_limits(self):
        (self.root / "linked.csv").symlink_to(self.root / "base.csv")
        for path in ("linked.csv", "../base.csv", ".git/config"):
            with self.subTest(path=path), self.assertRaises((OSError, ValueError)):
                capture_source(self.root, path)
        with self.assertRaises(ValueError):
            capture_source(self.root, "base.csv", max_bytes=3)

    def test_key_mapping_reorders_by_real_key_not_positional_ids(self):
        base = parse_source(capture_source(self.root, "base.csv"), ImportOptions())
        remote = parse_source(capture_source(self.root, "data.csv"), ImportOptions())
        local = read_csv_text(self.original.decode())
        mapping = (ColumnMapping("col_1", "col_1", "col_2"),
                   ColumnMapping("col_2", "col_2", "col_1"))
        before = deepcopy(local.to_content_dict())
        mapped_base, mapped_remote = map_source_tables(base, remote, local, mapping, key_column="col_1")
        self.assertEqual(mapped_base.to_content_dict(), local.to_content_dict())
        self.assertEqual([r.id for r in mapped_remote.rows[:3]], ["row_000", "row_002", "row_001"])
        self.assertNotIn(mapped_remote.rows[3].id, [r.id for r in local.rows])
        self.assertEqual(mapped_remote.cell("row_001", "col_2").value, 11)
        self.assertEqual(mapped_remote.cell("row_002", "col_2").value, 21)
        self.assertEqual(local.to_content_dict(), before)

    def test_ambiguous_key_or_incomplete_column_mapping_is_refused(self):
        local = read_csv_text(self.original.decode())
        base = parse_source(capture_source(self.root, "base.csv"), ImportOptions())
        mapping = (ColumnMapping("col_1", "col_1", "col_1"), ColumnMapping("col_2", "col_2", "col_2"))
        (self.root / "data.csv").write_bytes(b"Key,Value\nA,2\nA,3\n")
        remote = parse_source(capture_source(self.root, "data.csv"), ImportOptions())
        for mappings in (mapping, mapping[:1]):
            with self.subTest(mappings=mappings), self.assertRaises(ValueError):
                map_source_tables(base, remote, local, mappings, key_column="col_1")

    def test_raw_text_keys_are_not_stripped_or_converted(self):
        (self.root / "base.csv").write_bytes(b"Key,Value\n 001 ,10\n")
        data = parse_source(capture_source(self.root, "base.csv"), ImportOptions())
        self.assertEqual(data.rows[1].cells["col_1"].value, " 001 ")
        self.assertEqual(data.rows[1].cells["col_1"].kind, "text")

    def test_explicit_encoding_rejects_invalid_bytes_without_replacement(self):
        (self.root / "bad.csv").write_bytes(b"Key\n\xff\n")
        captured = capture_source(self.root, "bad.csv")
        with self.assertRaises(UnicodeError):
            parse_source(captured, ImportOptions())
        data = parse_source(captured, replace(ImportOptions(), encoding="latin-1"))
        self.assertEqual(data.rows[1].cells["col_1"].value, "ÿ")

    def test_xlsx_underreported_dimensions_do_not_hide_body_cells(self):
        import openpyxl
        book = openpyxl.Workbook()
        book.active.append(["Key", "Value"])
        book.active.append(["A", 10])
        book.active.append(["B", 20])
        stream = io.BytesIO()
        book.save(stream)
        changed = io.BytesIO()
        with zipfile.ZipFile(stream) as source, zipfile.ZipFile(changed, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', data)
                target.writestr(item, data)
        (self.root / "data.xlsx").write_bytes(changed.getvalue())
        data = parse_source(capture_source(self.root, "data.xlsx"), replace(ImportOptions(), sheet="Sheet"))
        self.assertEqual(len(data.rows), 3)
        self.assertEqual(data.rows[2].cells["col_2"].value, 20)

    def test_xlsx_formula_without_cached_value_is_not_imported_as_empty(self):
        import openpyxl
        book = openpyxl.Workbook()
        book.active.append(["Key", "Value"])
        book.active.append(["A", "=1+1"])
        book.save(self.root / "data.xlsx")
        with self.assertRaises(ValueError):
            parse_source(capture_source(self.root, "data.xlsx"), replace(ImportOptions(), sheet="Sheet"))
