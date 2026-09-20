from __future__ import annotations

from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from app.core.blocks.table_import import (
    SpreadsheetImportAdapter,
    parse_clipboard,
    read_csv_bytes,
    read_csv_text,
)


class CsvAdapterTests(TestCase):
    def test_empty_rows_are_not_silently_removed(self):
        data = read_csv_text("1\n\n2\n", header_row=False)
        self.assertEqual(len(data.rows), 3)
        self.assertEqual(data.cell("row_002", "col_1").kind, "empty")
        self.assertEqual(data.cell("row_003", "col_1").value, 2)

    def test_imported_header_and_first_body_row_both_reach_renderer(self):
        from app.core.blocks.table_renderer import render_table
        data = read_csv_text("Quantity,Note\n20,first row\n30,second row\n")
        self.assertEqual(data.rows[0].id, "row_000")
        self.assertEqual(data.cell("row_000", "col_1").value, "Quantity")
        self.assertEqual(data.cell("row_001", "col_1").value, 20)
        self.assertEqual(data.validate(), [])
        latex = render_table(data)
        self.assertIn("Quantity", latex)
        self.assertIn("first row", latex)
        self.assertIn("second row", latex)

    def test_limits_reject_instead_of_returning_partial_table(self):
        with patch("app.core.blocks.table_import.MAX_ROWS", 2), patch("app.core.blocks.table_import.MAX_COLS", 2):
            for text in ("a\n1\n2\n3\n", "a,b,c\n1,2,3\n"):
                with self.subTest(text=text), self.assertRaises(ValueError):
                    read_csv_text(text)
            self.assertEqual(len(read_csv_text("a\n1\n2\n").rows), 3)  # Header plus both body rows.

    def test_extra_cells_beyond_short_header_are_preserved(self):
        data = read_csv_text("a,b\n1,2,3\n")
        self.assertEqual(len(data.columns), 3)
        self.assertEqual(data.cell("row_001", "col_3").value, 3)

    def test_rfc4180_quotes_and_multiline_fields(self) -> None:
        text = 'name,note\n"Alice","line1\nline2"\n"Bob","said ""hi"""\n'
        data = read_csv_text(text)
        self.assertEqual(data.column_ids(), ["col_1", "col_2"])
        self.assertEqual(data.cell("row_001", "col_2").value, "line1\nline2")
        self.assertEqual(data.cell("row_002", "col_2").value, 'said "hi"')

    def test_type_inference_and_empty_values(self) -> None:
        text = "温度,备注\n20,基准\n1.5,\n,缺失\n"
        data = read_csv_text(text)
        self.assertEqual(data.columns[0].dataType, "number")
        self.assertEqual(data.columns[1].dataType, "text")
        self.assertEqual(data.cell("row_001", "col_1").value, 20)
        self.assertEqual(data.cell("row_002", "col_2").kind, "empty")

    def test_ragged_rows_are_padded(self) -> None:
        text = "a,b,c\n1,2\n3,4,5,6\n"
        data = read_csv_text(text)
        self.assertEqual(data.cell("row_001", "col_3").kind, "empty")
        self.assertEqual(data.cell("row_002", "col_3").value, 5)
        self.assertEqual(data.cell("row_002", "col_3").kind, "number")
        self.assertEqual(data.cell("row_002", "col_4").value, 6)

    def test_encoding_detection(self) -> None:
        utf8_bom = "温度\n20\n".encode("utf-8-sig")
        self.assertEqual(read_csv_bytes(utf8_bom).columns[0].name, "温度")
        gbk = "温度\n20\n".encode("gbk")
        self.assertEqual(read_csv_bytes(gbk).columns[0].name, "温度")

    def test_oversized_source_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            read_csv_bytes(b"x" * 100, max_bytes=10)


class ClipboardAdapterTests(TestCase):
    def test_tsv_rectangle(self) -> None:
        text = "A\tB\n1\t2\n3\t4\n"
        data = parse_clipboard(text)
        assert data is not None
        self.assertEqual(len(data.rows), 3)
        self.assertEqual(data.rows[2].cells["col_2"].value, 4)

    def test_html_table(self) -> None:
        text = "<table><tr><td>1</td><td>2</td></tr><tr><td>3</td><td>4</td></tr></table>"
        data = parse_clipboard(text)
        assert data is not None
        self.assertEqual(len(data.rows), 2)
        self.assertEqual(data.rows[0].cells["col_2"].value, 2)


class SpreadsheetAdapterTests(TestCase):
    def test_sheet_and_merge_metadata_use_the_same_captured_workbook_bytes(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "captured.xlsx"
            import openpyxl
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["before", "sibling"])
            workbook.save(path)
            with SpreadsheetImportAdapter(path) as adapter:
                sheet["A1"] = "after"
                sheet.merge_cells("A1:B1")
                workbook.save(path)
                data = adapter.read_sheet(sheet.title, header_row=False)
                self.assertEqual(data.cell("row_001", "col_1").value, "before")
                self.assertEqual(data.cell("row_001", "col_2").value, "sibling")
                self.assertEqual(data.merges, [])

    def test_expanded_archive_limit_is_checked_before_openpyxl(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "expanded.xlsx"
            import zipfile
            with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("oversize.xml", b"x" * 20000)
            with patch("openpyxl.load_workbook") as load, self.assertRaises(ValueError):
                SpreadsheetImportAdapter(path, max_bytes=2000)
            load.assert_not_called()

    def test_merge_coordinates_and_range_crop_follow_table_row_order(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "merged.xlsx"
            import openpyxl
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            for row in (("a", "b", "c", "d"), (1, 2, 3, 4), (5, 6, 7, 8), (9, 10, 11, 12)):
                sheet.append(row)
            sheet.merge_cells("B2:C3")
            workbook.save(path)
            with SpreadsheetImportAdapter(path) as adapter:
                data = adapter.read_sheet(sheet.title)
                self.assertEqual(data.merges, [[1, 2, 1, 2]])
                self.assertEqual(data.validate(), [])
                cropped = adapter.read_sheet(sheet.title, cell_range="B2:C3", header_row=False)
                self.assertEqual(cropped.merges, [[0, 1, 0, 1]])
                self.assertEqual(cropped.validate(), [])
                with self.assertRaises(ValueError):
                    adapter.read_sheet(sheet.title, cell_range="A1:B2", header_row=False)

    def test_row_column_and_range_limits_reject_without_truncation(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "limits.xlsx"
            import openpyxl
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["a", "b", "c"])
            sheet.append([1, 2, 3])
            sheet.append([4, 5, 6])
            workbook.save(path)
            for limits in ({"max_rows": 1}, {"max_cols": 2}):
                with self.subTest(limits=limits), SpreadsheetImportAdapter(path, **limits) as adapter:
                    with self.assertRaises(ValueError):
                        adapter.read_sheet(sheet.title)
            with SpreadsheetImportAdapter(path, max_rows=2, max_cols=2) as adapter:
                result = adapter.read_sheet(sheet.title, cell_range="B2:C3", header_row=False)
                self.assertEqual(len(result.rows), 2)
                self.assertEqual(result.cell("row_002", "col_2").value, 6)
                with self.assertRaises(ValueError):
                    adapter.read_sheet(sheet.title, cell_range="A0:B1")

    def test_multi_sheet_range_header_dates_and_merges(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "book.xlsx"
            import openpyxl

            workbook = openpyxl.Workbook()
            first = workbook.active
            first.title = "Experiment 1"
            first.append(["温度", "反应速率"])
            first.append([20, 1.25])
            first.append([datetime(2024, 1, 2), None])
            first.merge_cells("A1:B1")
            second = workbook.create_sheet("Notes")
            second.append(["x", "y"])
            workbook.save(str(path))

            adapter = SpreadsheetImportAdapter(path)
            self.assertEqual(adapter.sheet_names(), ["Experiment 1", "Notes"])
            data = adapter.read_sheet("Experiment 1")
            self.assertEqual(data.columns[0].name, "温度")
            # 数字与日期混合 -> 保守推断为 text，具体单元格保留各自类型。
            self.assertEqual(data.columns[0].dataType, "text")
            self.assertEqual(data.cell("row_001", "col_1").kind, "number")
            self.assertEqual(data.cell("row_002", "col_1").kind, "date")
            self.assertEqual(len(data.merges), 1)

            ranged = adapter.read_sheet("Experiment 1", cell_range="A1:B2", header_row=True)
            self.assertEqual(len(ranged.rows), 2)  # Preserve the selected header and data row.
            self.assertEqual(ranged.cell("row_001", "col_2").value, 1.25)

            notes = adapter.read_sheet("Notes", header_row=False)
            self.assertEqual(notes.rows[0].cells["col_1"].value, "x")

    def test_formula_cells_are_not_executed(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "formula.xlsx"
            import openpyxl

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = 2
            sheet["C1"] = "=A1+B1"
            workbook.save(str(path))

            # An uncalculated formula is neither evaluated nor silently
            # converted into an accepted empty cell. No table is returned.
            with SpreadsheetImportAdapter(path) as adapter, self.assertRaisesRegex(ValueError, "no cached value"):
                adapter.read_sheet(workbook.active.title, header_row=False)

    def test_unsupported_sheet_raises(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "book.xlsx"
            import openpyxl

            workbook = openpyxl.Workbook()
            workbook.save(str(path))
            with self.assertRaises(ValueError):
                SpreadsheetImportAdapter(path).read_sheet("Nope")

    def test_oversized_workbook_is_rejected(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "book.xlsx"
            path.write_bytes(b"x" * 100)
            with self.assertRaises(ValueError):
                SpreadsheetImportAdapter(path, max_bytes=10)
