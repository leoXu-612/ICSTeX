"""Synthetic merge candidate, import fidelity and explicit table FINAL evidence.

This is not linked-source application/CAS/Undo or native IME/AX acceptance.
All fixture writes and output stay under a fresh --output directory.
"""
from copy import deepcopy
import argparse
import hashlib
import json
from pathlib import Path
import platform
import sys

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

from PySide6.QtCore import QEvent, QPoint, QSettings, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QDialog, QDialogButtonBox, QScrollArea
from app.core.blocks.source_merge import merge_three_way
from app.core.blocks.table_import import read_csv_text, SpreadsheetImportAdapter
from app.core.blocks.table_model import Cell
from app.core.blocks.table_renderer import render_table, required_packages
from app.core.latex_tools import LaTeXEngine
from app.core.pdf_state import PdfFreshness
from app.core.settings import AppSettings
from app.gui.blocks.merge_dialog import MergeDialog
from app.gui.main_window import MainWindow
from app.gui.theme import apply_theme
from tools.probe_submission_check import wait_for


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    output = parser.parse_args().output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    application = QApplication.instance() or QApplication([])
    apply_theme(application)
    digest = hashlib.sha256()
    for path in sorted((REPO / "app").rglob("*.py")):
        digest.update(path.relative_to(REPO).as_posix().encode() + b"\0" + path.read_bytes() + b"\0")
    evidence = {"app_python_tree_sha256": digest.hexdigest(), "platform": platform.platform(),
                "python": platform.python_version(), "qt_platform": application.platformName()}
    base = read_csv_text("Key,Value,Note,Flag\nA,10,untouched,false\nB,20,second,true\n")
    base.table_id = "synthetic-table"
    remote, local = deepcopy(base), deepcopy(base)
    remote.rows[1].cells.update(col_2=Cell("number", 11), col_3=Cell("text", "remote note"))
    local.rows[1].cells.update(col_2=Cell("number", 12), col_3=Cell("text", "local note"))
    remote.rows[2].cells["col_2"] = Cell("number", 21)
    originals = [data.to_content_dict() for data in (base, remote, local)]
    result = merge_three_way(base, remote, local)
    assert len(result.conflicts) == 2
    window = MainWindow(settings_store=AppSettings(QSettings(str(output / "qa.ini"), QSettings.Format.IniFormat)))
    window.auto_compile_action.setChecked(False)
    dialog = MergeDialog(result, window)
    try:
        dialog.show()
        application.processEvents()
        scroll = dialog.findChild(QScrollArea)
        scroll.ensureWidgetVisible(dialog._remote_buttons[0])
        application.processEvents()
        button = dialog._remote_buttons[0]
        QTest.mouseClick(button, Qt.MouseButton.LeftButton, pos=QPoint(10, button.height() // 2))
        assert button.isChecked()
        scroll.ensureWidgetVisible(dialog._manual_edits[1])
        dialog._manual_edits[1].setFocus()
        QTest.keyClicks(dialog._manual_edits[1], "  reviewed note  ")
        ok = dialog.buttons.button(QDialogButtonBox.StandardButton.Ok)
        assert not ok.isEnabled()
        dialog.preview_button.click()
        assert ok.isEnabled()
        assert "reviewed note" in dialog.preview.toPlainText()
        evidence["scale_checks"] = []
        for scale in (0.9, 1.0, 1.1, 1.25, 1.5):
            window.set_ui_scale(scale)
            dialog.resize(760, 640)
            application.processEvents()
            for button in (dialog.preview_button, ok):
                assert dialog.rect().contains(button.mapTo(dialog, button.rect().center()))
            assert dialog.grab().save(str(output / f"merge-preview-{round(scale * 100)}.png"))
            evidence["scale_checks"].append([scale, dialog.width(), dialog.height()])
        ok.click()
        assert QDialog.result(dialog) == QDialog.DialogCode.Accepted
        candidate = dialog.result.data
        assert [row.id for row in candidate.rows] == ["row_000", "row_001", "row_002"]
        assert candidate.cell("row_001", "col_1").value == "A"
        assert candidate.cell("row_001", "col_2").value == 11
        assert candidate.cell("row_001", "col_3").value == "  reviewed note  "
        assert candidate.cell("row_001", "col_4").value is False
        assert candidate.cell("row_002", "col_2").value == 21
        assert [data.to_content_dict() for data in (base, remote, local)] == originals
        assert candidate.table_id == "synthetic-table"
        evidence["same_row_choices_preserve_siblings_false_whitespace_and_inputs"] = True
        cancelled = MergeDialog(result, window)
        cancelled._remote_buttons[0].setChecked(True)
        cancelled.preview_button.click()
        cancelled.reject()
        assert cancelled.result.data.to_content_dict() == result.data.to_content_dict()
        cancelled.deleteLater()
        evidence["cancel_does_not_apply_candidate"] = True

        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Quantity", "Note"])
        sheet.append([20, "first body"])
        sheet.append([30, "second body"])
        path = output / "synthetic-import.xlsx"
        workbook.save(path)
        raw = path.read_bytes()
        with SpreadsheetImportAdapter(path) as adapter:
            imported = adapter.read_sheet(sheet.title)
            assert len(imported.rows) == 3
            assert imported.cell("row_000", "col_1").value == "Quantity"
            assert imported.cell("row_001", "col_1").value == 20
        assert path.read_bytes() == raw
        for input_text in ("x\n" + "1\n" * 5001, ",".join(str(i) for i in range(201))):
            try:
                read_csv_text(input_text)
            except ValueError:
                pass
            else:
                raise AssertionError("Oversize CSV returned a partial table")
        evidence["xlsx_header_body_preserved_and_csv_limits_refused"] = True

        project = output / "synthetic-project"
        project.mkdir()
        root = project / "main.tex"
        packages = set(required_packages(candidate)) | set(required_packages(imported))
        source = "\\documentclass{article}\n" + "".join(f"\\usepackage{{{name}}}\n" for name in sorted(packages))
        source += "\\begin{document}\n" + render_table(candidate, caption="Resolved candidate")
        source += render_table(imported, caption="Imported header and body") + "\\end{document}\n"
        root.write_text(source)
        before = root.read_bytes()
        window.set_ui_scale(1.0)
        window.open_file(root)
        window.compile.set_engine(LaTeXEngine.PDFLATEX, compile_after=False)
        window.resize(1120, 820)
        window.show()
        assert not window.compile_authorized_roots
        window.compile_action.trigger()
        wait_for(lambda: window.pdf_state.record_for(root).freshness == PdfFreshness.CURRENT)
        wait_for(lambda: window.pdf_panel._document is not None and window.pdf_panel._document.pageCount() > 0)
        record = window.pdf_state.record_for(root)
        payload = record.last_successful_pdf.read_bytes()
        (output / "merge-import-final.pdf").write_bytes(payload)
        assert root.read_bytes() == before
        evidence["explicit_final_pdf_sha256"] = hashlib.sha256(payload).hexdigest()
        evidence["explicit_final_pdf_pages"] = window.pdf_panel._document.pageCount()
        assert window.grab().save(str(output / "merge-import-final-window.png"))
    finally:
        dialog.close()
        dialog.deleteLater()
        for tab in window.tabs.values():
            window.documents.cancel_save_timer(tab)
            tab.modified = tab.dirty = False
        window.close()
        window.deleteLater()
        application.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    evidence["limits"] = ["Candidate confirmation does not apply to ProjectSession or update SourceRecord",
        "Genuine source baseline, mapping, target/source CAS and one Undo remain integration work",
        "Synthetic offscreen events do not prove native keyboard/IME/AX/Windows acceptance",
        "No original project, package, installed application, commit or release mutation"]
    (output / "report.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(evidence, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
