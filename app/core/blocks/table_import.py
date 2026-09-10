"""Table import adapters (Sprint 3): CSV, clipboard, and XLSX.

All external files are untrusted: spreadsheets are read with cached values
only (formulas are never evaluated, macros are never executed), sizes are
bounded, and third-party reading is wrapped behind adapter boundaries.
"""
from __future__ import annotations

import csv
from datetime import date, datetime
from html.parser import HTMLParser
import html as html_lib
import io
import os
import posixpath
from pathlib import Path
import re
import zipfile
from xml.etree import ElementTree

from app.core.blocks.table_model import Cell, ColumnSpec, TableData, TableRow
from app.core.table_clipboard import MAX_CLIPBOARD_CELLS, parse_grid


ENCODINGS = ("utf-8-sig", "utf-8", "gbk", "latin-1")
NUMBER_RE = re.compile(r"^[+-]?(\d+([.]\d*)?|[.]\d+)([eE][+-]?\d+)?$")
DATE_RE = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})$")
MAX_ROWS = 5000
MAX_COLS = 200
MAX_SOURCE_BYTES = 200 * 1024 * 1024

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def detect_encoding(data: bytes) -> str:
    for encoding in ENCODINGS:
        try:
            data.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "latin-1"


def infer_cell(value: object) -> Cell:
    if value is None or value == "":
        return Cell(kind="empty")
    if isinstance(value, bool):
        return Cell(kind="boolean", value=bool(value))
    if isinstance(value, (int, float)):
        return Cell(kind="number", value=value)
    if isinstance(value, (datetime, date)):
        return Cell(kind="date", value=value.isoformat())
    text = str(value).strip()
    if not text:
        return Cell(kind="empty")
    if text.lower() in ("true", "false"):
        return Cell(kind="boolean", value=text.lower() == "true")
    if NUMBER_RE.fullmatch(text):
        return Cell(kind="number", value=_to_number(text))
    match = DATE_RE.fullmatch(text)
    if match:
        return Cell(kind="date", value=f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}")
    return Cell(kind="text", value=text)


def _to_number(text: str) -> object:
    try:
        return int(text)
    except ValueError:
        return float(text)


def column_spec_from_cells(column_id: str, name: str, cells: list[Cell]) -> ColumnSpec:
    kinds = [cell.kind for cell in cells if cell.kind != "empty"]
    if kinds and all(kind == "number" for kind in kinds):
        data_type = "number"
        alignment = "decimal"
    elif kinds and all(kind == "boolean" for kind in kinds):
        data_type = "boolean"
        alignment = "center"
    elif kinds and all(kind == "date" for kind in kinds):
        data_type = "date"
        alignment = "center"
    else:
        data_type = "text"
        alignment = "left"
    return ColumnSpec(id=column_id, name=name, dataType=data_type, alignment=alignment)


def _delimiter(text: str) -> str:
    first = next((line for line in text.splitlines() if line.strip()), "")
    if "\t" in first and first.count("\t") >= first.count(","):
        return "\t"
    return ","


def read_csv_text(text: str, *, header_row: bool = True, infer_types: bool = True) -> TableData:
    if len(text) > MAX_SOURCE_BYTES or len(text.encode("utf-8")) > MAX_SOURCE_BYTES:
        raise ValueError("CSV text exceeds the source byte limit.")
    reader = csv.reader(io.StringIO(text), delimiter=_delimiter(text), strict=True)
    raw_rows = []
    for row in reader:
        if len(row) > MAX_COLS or len(raw_rows) >= MAX_ROWS + int(header_row):
            raise ValueError("CSV row/column limit exceeded; no partial table was imported.")
        raw_rows.append(row)
    return _table_from_grid(raw_rows, header_row=header_row, infer_types=infer_types)


def read_csv_bytes(
    data: bytes,
    *,
    header_row: bool = True,
    max_bytes: int = MAX_SOURCE_BYTES,
) -> TableData:
    if len(data) > max_bytes:
        raise ValueError(f"文件过大（>{max_bytes} 字节），已拒绝导入。")
    encoding = detect_encoding(data)
    return read_csv_text(data.decode(encoding), header_row=header_row)


def parse_clipboard(text: str) -> TableData | None:
    if not text.strip():
        return None
    lowered = text.lower()
    if "<table" in lowered:
        rows = _parse_html_table(text)
        if rows:
            return _table_from_grid(rows, header_row=False)
        return None
    return read_csv_text(text, header_row=False)


def parse_clipboard_grid(text: str, *, max_rows: int, max_columns: int) -> list[list[str]]:
    """Editor paste keeps HTML support, but never silently clips a rectangle."""
    if "<table" not in text.lower():
        return parse_grid(text, max_rows=max_rows, max_columns=max_columns)
    if len(text) > 2_000_000:
        raise ValueError("Clipboard text is too large.")
    rows = _parse_html_table(text)
    width = max((len(row) for row in rows), default=0)
    if len(rows) > max_rows or width > max_columns or len(rows) * width > MAX_CLIPBOARD_CELLS:
        raise ValueError("Clipboard rectangle exceeds editor limits.")
    return [row + [""] * (width - len(row)) for row in rows]


def _table_from_grid(raw_rows: list[list[str]], *, header_row: bool, infer_types: bool = True) -> TableData:
    if not raw_rows:
        return TableData()
    width = max(len(row) for row in raw_rows)
    if width > MAX_COLS or len(raw_rows) > MAX_ROWS + int(header_row):
        raise ValueError("Table row/column limit exceeded; no partial table was imported.")
    if header_row:
        names = list(raw_rows[0]) + [""] * (width - len(raw_rows[0]))
        header = [name.strip() or f"Col{index + 1}" for index, name in enumerate(names)]
        header = _unique_names(header)
        body = raw_rows[1:]
    else:
        header = [f"Col{index + 1}" for index in range(width)]
        body = raw_rows

    columns = [
        ColumnSpec(
            id=f"col_{index + 1}",
            name=name,
            dataType="text",
            alignment="left",
        )
        for index, name in enumerate(header)
    ]
    cells_grid: list[list[Cell]] = []
    for raw in body:
        padded = list(raw[:width])
        while len(padded) < width:
            padded.append("")
        cells_grid.append([infer_cell(cell) if infer_types else Cell("text", cell) for cell in padded])
    for column_index, column in enumerate(columns):
        values = [row[column_index] for row in cells_grid]
        inferred = column_spec_from_cells(column.id, column.name, values)
        columns[column_index] = inferred
    rows = [
        TableRow(
            id=f"row_{index + 1:03d}",
            cells={column.id: cells_grid[index][column_index] for column_index, column in enumerate(columns)},
        )
        for index in range(len(cells_grid))
    ]
    if header_row:
        rows.insert(0, TableRow("row_000", {column.id: Cell("text", names[index])
                                           for index, column in enumerate(columns)}))
    return TableData(columns=columns, rows=rows, header_row_count=1 if header_row else 0)


def _unique_names(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for name in names:
        if name in seen:
            seen[name] += 1
            result.append(f"{name}_{seen[name]}")
        else:
            seen[name] = 0
            result.append(name)
    return result


class _HTMLTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._in_cell = False
        self._current_cell: list[str] = []
        self._current_row: list[str] = []
        self.rows: list[list[str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("td", "th"):
            self._in_cell = True
            self._current_cell = []
        elif tag == "tr":
            self._current_row = []

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th"):
            self._in_cell = False
            self._current_row.append(" ".join(self._current_cell).strip())
        elif tag == "tr":
            if self._current_row:
                self.rows.append(self._current_row)

    def handle_data(self, data: str) -> None:
        if self._in_cell:
            self._current_cell.append(data)


def _parse_html_table(text: str) -> list[list[str]]:
    parser = _HTMLTableParser()
    parser.feed(text)
    return parser.rows


class SpreadsheetImportAdapter:
    """Reads XLSX workbooks with cached values only (no formula execution)."""

    def __init__(
        self,
        path: Path,
        *,
        max_rows: int = MAX_ROWS,
        max_cols: int = MAX_COLS,
        max_bytes: int = MAX_SOURCE_BYTES,
    ) -> None:
        self._path = Path(path)
        with self._path.open("rb") as stream:
            before = os.fstat(stream.fileno())
            if before.st_size > max_bytes:
                raise ValueError(f"文件过大（>{max_bytes} 字节），已拒绝导入。")
            self._bytes = stream.read(max_bytes + 1)
            after = os.fstat(stream.fileno())
        if len(self._bytes) > max_bytes or (before.st_size, before.st_mtime_ns, before.st_ctime_ns) != (
                after.st_size, after.st_mtime_ns, after.st_ctime_ns):
            raise ValueError("Workbook changed or exceeded the byte limit while reading.")
        self._load_bytes(self._bytes, max_rows=max_rows, max_cols=max_cols, max_bytes=max_bytes)

    @classmethod
    def from_bytes(cls, data: bytes, *, max_rows=MAX_ROWS, max_cols=MAX_COLS, max_bytes=MAX_SOURCE_BYTES):
        """Parse an already authorized immutable capture without reopening a path."""
        instance = cls.__new__(cls)
        instance._load_bytes(data, max_rows=max_rows, max_cols=max_cols, max_bytes=max_bytes)
        return instance

    def _load_bytes(self, data, *, max_rows, max_cols, max_bytes):
        import openpyxl

        if len(data) > max_bytes:
            raise ValueError("Workbook exceeds the source byte limit.")
        self._bytes = bytes(data)
        with zipfile.ZipFile(io.BytesIO(self._bytes)) as archive:
            members = archive.infolist()
            if (len(members) > 10000 or len({item.filename for item in members}) != len(members)
                    or sum(item.file_size for item in members) > max_bytes):
                raise ValueError("Expanded workbook exceeds the archive limit.")
        self._workbook = openpyxl.load_workbook(io.BytesIO(self._bytes), read_only=True, data_only=True, keep_links=False)
        self._max_rows = max_rows
        self._max_cols = max_cols

    def close(self) -> None:
        self._workbook.close()

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        self.close()

    def sheet_names(self) -> list[str]:
        return list(self._workbook.sheetnames)

    def read_sheet(self, name: str, *, cell_range: str | None = None, header_row: bool = True) -> TableData:
        if name not in self._workbook.sheetnames:
            raise ValueError(f"工作表不存在：{name}")
        sheet = self._workbook[name]
        sheet_xml = _sheet_xml(io.BytesIO(self._bytes), name)
        observed_row = observed_col = 0
        cells = []
        for row in sheet_xml.findall(f".//{{{_MAIN_NS}}}sheetData/{{{_MAIN_NS}}}row"):
            observed_row = max(observed_row, int(row.get("r", "0")))
            for cell in row.findall(f"{{{_MAIN_NS}}}c"):
                reference = cell.get("r", "")
                match = re.fullmatch(r"([A-Z]+)([1-9]\d*)", reference)
                if match is None:
                    raise ValueError("Worksheet cell has no explicit valid coordinate; import refused.")
                column, number = _column_letters_to_index(match.group(1)), int(match.group(2))
                observed_row, observed_col = max(observed_row, number), max(observed_col, column)
                cells.append((number, column, cell))
        min_row = min_col = 1
        if cell_range:
            min_col, min_row, max_col, max_row = _parse_range(cell_range)
            if max_row - min_row + 1 > self._max_rows + int(header_row) or max_col - min_col + 1 > self._max_cols:
                raise ValueError("Selected sheet range exceeds row/column limits; import refused.")
            merge_rows = min_row - 1
            merge_cols = min_col - 1
        else:
            max_row = max(sheet.max_row or 0, observed_row)
            max_col = max(sheet.max_column or 0, observed_col)
            if max_row > self._max_rows + int(header_row) or max_col > self._max_cols:
                raise ValueError("Worksheet exceeds row/column limits; no partial table was imported.")
            merge_rows = 0
            merge_cols = 0
        for row, column, cell in cells:
            if min_row <= row <= max_row and min_col <= column <= max_col:
                cached = cell.find(f"{{{_MAIN_NS}}}v")
                if cell.find(f"{{{_MAIN_NS}}}f") is not None and (cached is None or cached.text is None):
                    raise ValueError("Formula has no cached value; recalculate in the spreadsheet application before import.")
        # A dishonest/stale <dimension> must not hide actual cells. Explicit
        # observed/range bounds above remain authoritative for this import.
        sheet.reset_dimensions()
        values = list(sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col,
                                      max_col=max_col, values_only=True)) if max_row and max_col else []
        merges = []
        for r1, r2, c1, c2 in _merges_from_xml(sheet_xml):
            if r2 < min_row - 1 or r1 >= max_row or c2 < min_col - 1 or c1 >= max_col:
                continue
            if r1 < min_row - 1 or r2 >= max_row or c1 < min_col - 1 or c2 >= max_col:
                raise ValueError("Selected range cuts a merged cell; choose the complete merged region.")
            merges.append([r1 - merge_rows, r2 - merge_rows, c1 - merge_cols, c2 - merge_cols])
        grid = [[value for value in row] for row in values]
        data = _table_from_cells(grid, header_row=header_row)
        data.merges = merges
        return data


def _parse_range(cell_range: str) -> tuple[int, int, int, int]:
    match = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", cell_range.strip().upper())
    if match is None:
        raise ValueError(f"非法单元格范围：{cell_range}")
    min_col = _column_letters_to_index(match.group(1))
    min_row = int(match.group(2))
    max_col = _column_letters_to_index(match.group(3))
    max_row = int(match.group(4))
    if min_row < 1 or min_col < 1 or max_row < min_row or max_col < min_col:
        raise ValueError(f"范围顺序非法：{cell_range}")
    return min_col, min_row, max_col, max_row


def _column_letters_to_index(letters: str) -> int:
    result = 0
    for char in letters:
        result = result * 26 + (ord(char) - ord("A") + 1)
    return result


def _sheet_merges(path: Path, sheet_name: str) -> list[list[int]]:
    """Read mergeCell ranges directly from the worksheet XML (bounded)."""
    return _merges_from_xml(_sheet_xml(path, sheet_name))


def _sheet_xml(path, sheet_name):

    with zipfile.ZipFile(path) as archive:
        workbook_xml = archive.read("xl/workbook.xml")
        rels_xml = archive.read("xl/_rels/workbook.xml.rels")
    namespaces = {"m": _MAIN_NS}
    workbook_root = ElementTree.fromstring(workbook_xml)
    relationship_id: str | None = None
    for sheet in workbook_root.findall("m:sheets/m:sheet", namespaces):
        if sheet.get("name") == sheet_name:
            relationship_id = sheet.get(f"{{{_REL_NS}}}id")
            break
    if relationship_id is None:
        raise ValueError("Worksheet relationship is missing.")
    rel_root = ElementTree.fromstring(rels_xml)
    target: str | None = None
    for relationship in rel_root:
        if relationship.get("Id") == relationship_id:
            target = relationship.get("Target")
            break
    if target is None:
        raise ValueError("Worksheet archive target is missing.")
    sheet_path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
    if not sheet_path.startswith("xl/") or ".." in sheet_path.split("/"):
        raise ValueError("Invalid worksheet archive relationship.")
    with zipfile.ZipFile(path) as archive:
        sheet_xml = archive.read(sheet_path)
    return ElementTree.fromstring(sheet_xml)


def _merges_from_xml(sheet_root):
    merges: list[list[int]] = []
    for cell in sheet_root.findall(f".//{{{_MAIN_NS}}}mergeCell"):
        parsed = _parse_cell_reference(str(cell.get("ref", "")))
        if parsed is not None:
            merges.append(parsed)
    return merges


def _parse_cell_reference(reference: str) -> list[int] | None:
    match = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", reference)
    if match is None:
        return None
    return [
        int(match.group(2)) - 1,
        int(match.group(4)) - 1,
        _column_letters_to_index(match.group(1)) - 1,
        _column_letters_to_index(match.group(3)) - 1,
    ]


def _table_from_cells(grid: list[list[object]], *, header_row: bool) -> TableData:
    if not grid:
        return TableData()
    width = max(len(row) for row in grid)
    if width > MAX_COLS or len(grid) > MAX_ROWS + int(header_row):
        raise ValueError("Table row/column limit exceeded; no partial table was imported.")
    if header_row:
        header_values = list(grid[0]) + [None] * (width - len(grid[0]))
        names = _unique_names(
            [str(cell).strip() if cell is not None else f"Col{index + 1}" for index, cell in enumerate(header_values)]
        )
        body = grid[1:]
    else:
        names = [f"Col{index + 1}" for index in range(width)]
        body = grid
    columns = [ColumnSpec(id=f"col_{index + 1}", name=name) for index, name in enumerate(names)]
    cells_grid: list[list[Cell]] = []
    for raw in body:
        padded = list(raw[:width])
        while len(padded) < width:
            padded.append(None)
        cells_grid.append([infer_cell(cell) for cell in padded])
    for column_index, column in enumerate(columns):
        values = [row[column_index] for row in cells_grid]
        columns[column_index] = column_spec_from_cells(column.id, column.name, values)
    rows = [
        TableRow(
            id=f"row_{index + 1:03d}",
            cells={column.id: cells_grid[index][column_index] for column_index, column in enumerate(columns)},
        )
        for index in range(len(cells_grid))
    ]
    if header_row:
        rows.insert(0, TableRow("row_000", {column.id: Cell("text", "" if header_values[index] is None
                                                          else str(header_values[index]))
                                           for index, column in enumerate(columns)}))
    return TableData(columns=columns, rows=rows, header_row_count=1 if header_row else 0)
